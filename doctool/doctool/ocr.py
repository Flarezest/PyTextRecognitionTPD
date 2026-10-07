"""Загрузка документов и OCR-бэкенды (Tesseract + опциональная локальная VLM через Ollama)."""
from __future__ import annotations

import base64
import json
import os
import re
import shutil
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path

import cv2
import numpy as np
import pymupdf
import pytesseract
from rapidfuzz import fuzz

# путь к tesseract: переменная TESSERACT_CMD, PATH или стандартная папка Windows
_tess = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
if not _tess and os.name == "nt":
    for cand in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                 r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                 os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
                 os.path.expandvars(r"%LOCALAPPDATA%\Tesseract-OCR\tesseract.exe")):
        if os.path.exists(cand):
            _tess = cand
            break
if _tess:
    pytesseract.pytesseract.tesseract_cmd = _tess


def check_tesseract() -> str | None:
    """Возвращает текст ошибки, если Tesseract или нужные языки не установлены."""
    try:
        langs = set(pytesseract.get_languages(config=""))
    except (pytesseract.TesseractNotFoundError, OSError):
        return ("Tesseract не найден. Установите его (см. README, раздел «Установка») "
                "или укажите путь: set TESSERACT_CMD=C:\\путь\\к\\tesseract.exe")
    missing = {"rus", "eng", "osd"} - langs
    if missing:
        return (f"В Tesseract нет языков: {', '.join(sorted(missing))}. Скачайте файлы "
                f"<язык>.traineddata (github.com/tesseract-ocr/tessdata) в папку tessdata рядом с tesseract.exe. "
                f"Сейчас установлены: {', '.join(sorted(langs)) or 'нет'}")
    return None

# ---------------------------------------------------------------- загрузка

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}


@dataclass
class Page:
    image: np.ndarray | None          # BGR, None для чисто текстовых страниц
    text_layer: str = ""              # текстовый слой PDF (если есть)
    source: str = ""
    index: int = 0


def load_document(path: str | Path, dpi: int = 300) -> list[Page]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext in IMAGE_EXT:
        img = cv2.imdecode(np.fromfile(str(path), dtype=np.uint8), cv2.IMREAD_COLOR)
        return [Page(image=img, source=path.name, index=0)]
    if ext == ".pdf":
        pages = []
        with pymupdf.open(path) as doc:
            for i, p in enumerate(doc):
                pix = p.get_pixmap(dpi=dpi)
                img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
                img = cv2.cvtColor(img, cv2.COLOR_RGB2BGR if pix.n == 3 else cv2.COLOR_RGBA2BGR)
                pages.append(Page(image=img, text_layer=p.get_text(), source=path.name, index=i))
        return pages
    if ext == ".docx":
        from docx import Document  # python-docx
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [Page(image=None, text_layer="\n".join(parts), source=path.name)]
    raise ValueError(f"Неподдерживаемый формат: {path}")


# ---------------------------------------------------------------- Tesseract

@dataclass
class Word:
    text: str
    conf: float
    x: int
    y: int
    w: int
    h: int

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2


@dataclass
class Line:
    words: list[Word] = field(default_factory=list)

    @property
    def text(self) -> str:
        return " ".join(w.text for w in self.words)

    @property
    def conf(self) -> float:
        return float(np.mean([w.conf for w in self.words])) if self.words else 0.0

    @property
    def box(self) -> tuple[int, int, int, int]:
        x0 = min(w.x for w in self.words); y0 = min(w.y for w in self.words)
        x1 = max(w.x + w.w for w in self.words); y1 = max(w.y + w.h for w in self.words)
        return x0, y0, x1, y1


def tesseract_lines(img: np.ndarray, lang: str = "rus", psm: int = 3, config: str = "") -> list[Line]:
    data = pytesseract.image_to_data(img, lang=lang, config=f"--psm {psm} {config}",
                                     output_type=pytesseract.Output.DICT)
    lines: dict[tuple, Line] = {}
    for i, txt in enumerate(data["text"]):
        txt = (txt or "").strip()
        if not txt or float(data["conf"][i]) < 0:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, Line()).words.append(
            Word(txt, float(data["conf"][i]), data["left"][i], data["top"][i], data["width"][i], data["height"][i]))
    out = list(lines.values())
    out.sort(key=lambda l: (l.box[1], l.box[0]))
    return out


def tesseract_text(img: np.ndarray, lang: str = "rus", psm: int = 6, config: str = "") -> tuple[str, float]:
    lines = tesseract_lines(img, lang, psm, config)
    text = "\n".join(l.text for l in lines)
    conf = float(np.mean([w.conf for l in lines for w in l.words])) if lines else 0.0
    return text, conf


def rotate(img: np.ndarray, angle: int) -> np.ndarray:
    """Поворот на угол, кратный 90° (по часовой стрелке)."""
    return {0: img, 90: cv2.rotate(img, cv2.ROTATE_90_CLOCKWISE),
            180: cv2.rotate(img, cv2.ROTATE_180),
            270: cv2.rotate(img, cv2.ROTATE_90_COUNTERCLOCKWISE)}[angle % 360]


def downscale(img: np.ndarray, max_side: int) -> np.ndarray:
    h, w = img.shape[:2]
    s = max_side / max(h, w)
    return cv2.resize(img, None, fx=s, fy=s, interpolation=cv2.INTER_AREA) if s < 1 else img


def keyword_score(text: str, keywords: list[str]) -> int:
    """Сколько ключевых слов встречается в тексте. Нечётко: OCR часто рвёт и портит мелкие
    печатные подписи («ЛЕРАЦИЯ» вместо «ФЕДЕРАЦИЯ»)."""
    t = text.upper().replace("Ё", "Е")
    words = [w for w in re.findall(r"[А-ЯA-Z]{4,}", t)]
    score = 0
    for k in keywords:
        if k in t or any(fuzz.ratio(w, k) >= 75 or (len(w) >= 5 and fuzz.partial_ratio(w, k) >= 90) for w in words):
            score += 1
    return score


def _red_flat(img: np.ndarray) -> np.ndarray:
    ch = img[:, :, 2] if img.ndim == 3 else img
    small = cv2.resize(ch, None, fx=0.25, fy=0.25)
    bg = cv2.resize(cv2.medianBlur(small, 31), (ch.shape[1], ch.shape[0]))
    return cv2.divide(ch, bg, scale=255)


def best_orientation(img: np.ndarray, keywords: list[str], max_side: int = 1800) -> tuple[int, int]:
    """Выбирает поворот (0/90/180/270), при котором OCR находит больше ключевых слов.
    Сначала пробуется подсказка Tesseract OSD, затем остальные углы. Если по обычному изображению
    слов почти нет (рукописный паспорт, мелкие подписи на фоне сетки), второй проход идёт
    по «красному» каналу с выравниванием фона в разреженном режиме OCR."""
    small = downscale(img, max_side)
    order = [0, 90, 270, 180]
    try:
        osd = pytesseract.image_to_osd(small, config="--psm 0", output_type=pytesseract.Output.DICT)
        hint = int(osd.get("rotate", 0)) % 360
        order = [hint] + [a for a in order if a != hint]
    except pytesseract.TesseractError:
        pass
    best = (0, -1)
    enough = max(3, len(keywords) // 2)
    for a in order:
        txt, _ = tesseract_text(rotate(small, a), psm=3)
        sc = keyword_score(txt, keywords)
        if sc > best[1]:
            best = (a, sc)
        if sc >= enough:
            return best
    if best[1] < 3 and small.ndim == 3:
        flat = _red_flat(downscale(img, max(max_side, 2400)))
        for a in order:
            txt, _ = tesseract_text(rotate(flat, a), psm=11)
            sc = keyword_score(txt, keywords)
            if sc > best[1]:
                best = (a, sc)
    return best


def crop_to_content(img: np.ndarray, pad: int = 20) -> tuple[np.ndarray, bool]:
    """Обрезает белые поля скана (например, паспорт на половине листа A4).
    Возвращает исходник, если документ и так занимает почти весь кадр."""
    g = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY) if img.ndim == 3 else img
    small = downscale(g, 1200)
    k = g.shape[1] / small.shape[1]
    m = (small < 200).astype(np.uint8) * 255
    m = cv2.morphologyEx(m, cv2.MORPH_CLOSE, np.ones((15, 15), np.uint8))
    cnts, _ = cv2.findContours(m, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if not cnts:
        return img, False
    # объединяем все заметные области (у бледных сканов MRZ бывает отдельным «островом» под разворотом)
    total = small.shape[0] * small.shape[1]
    big = [c for c in cnts if cv2.contourArea(c) >= 0.003 * total] or [max(cnts, key=cv2.contourArea)]
    pts = np.vstack(big)
    x, y, w, h = cv2.boundingRect(pts)
    if w * h > 0.85 * small.shape[0] * small.shape[1] or w * h < 0.1 * small.shape[0] * small.shape[1]:
        return img, False
    x0, y0 = max(0, int(x * k) - pad), max(0, int(y * k) - pad)
    x1, y1 = min(g.shape[1], int((x + w) * k) + pad), min(g.shape[0], int((y + h) * k) + pad)
    return img[y0:y1, x0:x1], True


def remove_lines(gray: np.ndarray) -> np.ndarray:
    """Убирает длинные горизонтальные линии (подчёркивания, рамки таблиц): они мешают OCR
    и модели — например, линия под «22.05.2001» превращает «5» в «3»."""
    if gray.ndim == 3:
        gray = cv2.cvtColor(gray, cv2.COLOR_BGR2GRAY)
    bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    klen = max(40, gray.shape[1] // 12)
    lines = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (klen, 1)))
    lines = cv2.dilate(lines, np.ones((3, 3), np.uint8))
    out = gray.copy()
    out[lines > 0] = 255
    return out


def remove_lines_soft(gray: np.ndarray) -> np.ndarray:
    """Как remove_lines, но пиксели линии, через которые проходит штрих буквы (над или под линией есть
    чернила в том же столбце), остаются. Нужно для подчёркнутого печатного текста: remove_lines обрезает
    хвосты «g», «р», «у», «ф», и «trading» читается как «tradina»."""
    if gray.ndim == 3:
        gray = cv2.cvtColor(gray, cv2.COLOR_BGR2GRAY)
    bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    klen = max(40, gray.shape[1] // 12)
    lines = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (klen, 1)))
    lines = cv2.dilate(lines, np.ones((3, 1), np.uint8))
    ink = ((bw > 0) & (lines == 0)).astype(np.uint8)
    stroke = cv2.dilate(ink, np.ones((7, 1), np.uint8)) > 0
    out = gray.copy()
    out[(lines > 0) & ~stroke] = 255
    return out


# ---------------------------------------------------------------- VLM (Ollama)

class OllamaVLM:
    """Локальная vision-модель через Ollama (по умолчанию http://127.0.0.1:11434)."""

    def __init__(self, model: str = "qwen3-vl:4b-instruct", host: str = "http://127.0.0.1:11434",
                 timeout: int = 600, max_tokens: int = 64):
        self.model, self.host, self.timeout, self.max_tokens = model, host.rstrip("/"), timeout, max_tokens

    def available(self) -> bool:
        try:
            with urllib.request.urlopen(f"{self.host}/api/tags", timeout=3) as r:
                tags = json.load(r)
            return any(m.get("name", "").startswith(self.model.split(":")[0]) for m in tags.get("models", []))
        except Exception:
            return False

    def ask(self, img: np.ndarray, prompt: str, max_tokens: int | None = None) -> str:
        ok, buf = cv2.imencode(".png", img)
        payload = {
            "model": self.model, "prompt": prompt, "stream": False, "think": False,
            "images": [base64.b64encode(buf.tobytes()).decode()],
            "options": {"temperature": 0, "num_predict": max_tokens or self.max_tokens, "repeat_penalty": 1.2},
        }
        req = urllib.request.Request(f"{self.host}/api/generate", json.dumps(payload).encode(),
                                     {"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=self.timeout) as r:
            return json.load(r)["response"].strip()

    def transcribe(self, img: np.ndarray, hint: str = "") -> str:
        img = downscale(img, 1000)  # меньше пикселей — меньше токенов и быстрее на CPU
        prompt = ("Перепиши рукописный текст на изображении точно, как написано. "
                  "Печатный текст бланка (подписи в скобках, линии) не переписывай. "
                  "Если рукописного текста нет, ответь одним словом: ПУСТО. "
                  "Выведи только сам текст одной строкой, без пояснений.")
        out = self.ask(img, prompt)
        out = re.sub(r"^(вот|текст)[^:]*:\s*", "", out, flags=re.I).strip().strip('"«»')
        out = out.splitlines()[0] if out else ""
        # защита от «эха» и галлюцинаций на пустых полях
        if re.fullmatch(r"\W*(пусто|нет|empty|none)?\W*", out, flags=re.I) or re.fullmatch(r"[0\s_.\-]+", out):
            return ""
        if hint and fuzz.ratio(out.lower(), hint.lower()) > 80:
            return ""
        return out


# ---------------------------------------------------------------- фото документа

def _order_quad(pts: np.ndarray) -> np.ndarray:
    pts = pts.reshape(4, 2).astype("float32")
    s, d = pts.sum(1), np.diff(pts, axis=1).ravel()
    return np.array([pts[np.argmin(s)], pts[np.argmin(d)], pts[np.argmax(s)], pts[np.argmax(d)]], dtype="float32")


def straighten_photo(img: np.ndarray) -> tuple[np.ndarray, bool]:
    """Фото документа: ищет крупный четырёхугольник (лист/разворот) и выравнивает перспективу.
    Если уверенно найти не удалось — возвращает исходное изображение."""
    small = downscale(img, 1200)
    k = img.shape[1] / small.shape[1]
    gray = cv2.GaussianBlur(cv2.cvtColor(small, cv2.COLOR_BGR2GRAY), (5, 5), 0)
    edges = cv2.dilate(cv2.Canny(gray, 50, 150), np.ones((5, 5), np.uint8))
    cnts, _ = cv2.findContours(edges, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    area_img = small.shape[0] * small.shape[1]
    for c in sorted(cnts, key=cv2.contourArea, reverse=True)[:5]:
        a = cv2.contourArea(c)
        if a < 0.25 * area_img or a > 0.97 * area_img:
            continue
        approx = cv2.approxPolyDP(c, 0.02 * cv2.arcLength(c, True), True)
        if len(approx) != 4:
            continue
        q = _order_quad(approx) * k
        w = int(max(np.linalg.norm(q[0] - q[1]), np.linalg.norm(q[3] - q[2])))
        h = int(max(np.linalg.norm(q[0] - q[3]), np.linalg.norm(q[1] - q[2])))
        if w < 300 or h < 300 or not (0.3 < w / h < 3.5):
            continue
        dst = np.array([[0, 0], [w - 1, 0], [w - 1, h - 1], [0, h - 1]], dtype="float32")
        return cv2.warpPerspective(img, cv2.getPerspectiveTransform(q, dst), (w, h)), True
    return img, False
