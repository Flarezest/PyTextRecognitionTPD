"""Отчёты: JSON (для интеграции), Excel (реестр), HTML (ручная проверка с фрагментами)."""
from __future__ import annotations

import base64
import html
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .compare import STATUS_RU, Check, verdict
from .formspec import Fields
from .models import Extraction

FIELD_TITLES = {
    "applicant_header": "ФИО (шапка)", "applicant_header.birth_date": "Дата рождения",
    "applicant_header.raw_fio": "ФИО (как написано)", "inn": "ИНН ИП", "passport": "Серия, номер паспорта",
    "passport.issue_date": "Дата выдачи", "issued_by": "Кем выдан", "address": "Адрес регистрации",
    "applicant_fio": "ФИО («Я, …»)", "domains": "Домен(ы)", "domains.punycode": "Домен(ы), punycode",
    "domains.list": "Домен(ы), список", "domains.joined": "Домены, собранные из переносов строк",
    "domains.fragments": "Обрывки без доменной зоны", "domains.mixed": "Домены: смешаны кириллица и латиница",
    "domains.stated_count": "Доменов по тексту заявления («всего N»)",
    "domains.uncertain": "Домены, прочитанные по-разному",
    "services": "Доп. услуги", "new_admin_inline": "Новый администратор (текст)",
    "new_admin_org": "Новый администратор — юрлицо", "new_admin_fio": "Новый администратор — физлицо/ИП",
    "new_admin_contact": "Контакты нового администратора", "new_admin_contact.emails": "E-mail",
    "new_admin_contact.phones": "Телефон", "contract": "Договор/аккаунт", "signature_fio": "ФИО у подписи",
    "application_date": "Дата заявления", "application_date.raw": "Дата заявления (как написано)",
    "surname": "Фамилия", "given_name": "Имя", "patronymic": "Отчество", "sex": "Пол",
    "birth_date": "Дата рождения", "birth_place": "Место рождения", "series_number": "Серия и номер",
    "issue_date": "Дата выдачи", "department_code": "Код подразделения",
}


def field_title(key: str, ex: Extraction | None = None) -> str:
    """Название поля для отчёта: FIELD_TITLES (по имени поля или по роли) → title из YAML бланка → имя."""
    if key in FIELD_TITLES:
        return FIELD_TITLES[key]
    F = Fields(ex) if ex is not None else None
    base, dot, sub = key.partition(".")
    spec = ((F.form if F else None) or {}).get("fields", {}).get(base) or {}
    role = spec.get("role", base)
    if role + dot + sub in FIELD_TITLES:
        return FIELD_TITLES[role + dot + sub]
    if spec.get("title"):
        return spec["title"] + (f": {sub}" if dot else "")
    return key


COLORS = {"ok": "C6EFCE", "warn": "FFEB9C", "fail": "FFC7CE", "review": "DDEBF7", "info": "EDEDED"}


def _jsonable(ex: Extraction | None) -> dict | None:
    if ex is None:
        return None
    d = ex.to_dict()
    d["debug"] = {k: v for k, v in ex.debug.items() if k not in ("main_image",)}
    return d


def save_json(path: Path, app: Extraction, pas: Extraction | None, checks: list[Check]):
    data = {"created": datetime.now().isoformat(timespec="seconds"), "verdict": verdict(checks),
            "application": _jsonable(app), "passport": _jsonable(pas), "checks": [c.to_dict() for c in checks]}
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def append_excel(path: Path, case_id: str, app: Extraction, pas: Extraction | None, checks: list[Check],
                 decision=None):
    """Добавляет строку в реестр (один файл на все проверки) и лист с деталями."""
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Реестр"
        ws.append(["Дата проверки", "Дело", "Итог", "ФИО заявителя", "Паспорт", "Домен(ы)",
                   "Новый администратор", "Контакты", "Договор", "Расхождения", "Проверить вручную",
                   "Файл заявления", "Файл паспорта"])
        for c in ws[1]:
            c.font = Font(bold=True)
    ws = wb["Реестр"]
    g = Fields(app).get
    ws.append([datetime.now().strftime("%d.%m.%Y %H:%M"), case_id,
               (decision.title + ": " + "; ".join(decision.reasons)) if decision else verdict(checks),
               g("applicant_header") or g("applicant_fio") or "", g("passport") or "",
               g("domains") or "", g("new_admin_fio") or g("new_admin_org") or g("new_admin_inline") or "",
               g("new_admin_contact") or "", g("contract") or "",
               "; ".join(c.name for c in checks if c.status == "fail"),
               "; ".join(c.name for c in checks if c.status in ("review", "warn")),
               app.source_file, pas.source_file if pas else ""])
    res = ws.cell(ws.max_row, 3)
    if decision is not None:
        key = {"reject": "fail", "accept": "ok", "review": "warn"}[decision.code]
    else:
        key = ("fail" if "расхожд" in res.value.lower() and "не найдено" not in res.value
               else "ok" if "не найдено" in res.value else "warn")
    res.fill = PatternFill("solid", fgColor=COLORS[key])
    for col, w in zip("ABCDEFGHIJKLM", (16, 18, 24, 32, 14, 20, 32, 30, 22, 40, 40, 26, 26)):
        ws.column_dimensions[col].width = w

    name = case_id[:28]
    if name in wb.sheetnames:
        del wb[name]
    d = wb.create_sheet(name)
    d.append(["Проверка", "Статус", "Заявление", "Паспорт", "Комментарий"])
    for c in d[1]:
        c.font = Font(bold=True)
    for c in checks:
        d.append([c.name, STATUS_RU[c.status], c.application, c.passport, c.detail])
        d.cell(d.max_row, 2).fill = PatternFill("solid", fgColor=COLORS[c.status])
    d.append([])
    d.append(["Поле", "Значение", "Источник", "Уверенность", "Проверить"])
    for c in d[d.max_row]:
        c.font = Font(bold=True)
    for title, ex in (("Заявление", app), ("Паспорт", pas)):
        if ex is None:
            continue
        d.append([f"— {title} —"])
        for k, f in ex.fields.items():
            v = ", ".join(f.value) if isinstance(f.value, list) else f.value
            d.append([field_title(k, ex), str(v), f.source, f.confidence, "да" if f.needs_review else ""])
    for col, w in zip("ABCDE", (34, 40, 40, 40, 50)):
        d.column_dimensions[col].width = w
    for row in d.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def _img_tag(path: str | None, max_h: int = 120) -> str:
    if not path or not Path(path).exists():
        return ""
    b64 = base64.b64encode(Path(path).read_bytes()).decode()
    ext = Path(path).suffix.lstrip(".").replace("jpg", "jpeg")
    return f'<img src="data:image/{ext};base64,{b64}" style="max-height:{max_h}px;max-width:100%">'


def domains_table(domains: list[dict]) -> str:
    """Таблица «Домены в manager» для отчёта (domains — DomainInfo.to_dict()). Домены сгруппированы
    по проблемности (domains.PROBLEMS): сначала свободные по WHOIS, затем с несовпадающими данными и т. д."""
    from .domains import sort_by_problem
    e = html.escape
    st = {"found": "найден", "not_found": "не найден", "error": "ошибка"}
    out = []
    domains = sort_by_problem(domains)
    counts: dict[str, int] = {}
    for d in domains:
        counts[d.get("problem", "")] = counts.get(d.get("problem", ""), 0) + 1
    group = None
    for d in domains:
        if d.get("problem") and d.get("problem") != group:
            group = d["problem"]
            out.append(f'<tr class="grp"><td colspan="5">{e(d.get("problem_ru", group))} — {counts[group]}</td></tr>')
        cls = d.get("verdict") or ("fail" if d.get("status") == "not_found" else "review" if d.get("status") == "error" else "")
        if d.get("status") == "found":
            sd = d.get("sd") or {}
            data = (f"{e(d.get('holder', ''))}<br><span class=src>{e(d.get('kind_ru', ''))}"
                    + (f", паспорт {e(d.get('passport', ''))} от {e(sd.get('passport_date', ''))}" if d.get("kind") == "person" else "")
                    + (f", д. р. {e(sd.get('birth_date', ''))}" if sd.get("birth_date") else "")
                    + (f", ИНН {e(sd.get('code', ''))}" if sd.get("code") else "")
                    + (f"<br>e-mail: {e(', '.join(d['emails']))}" if d.get("emails") else "") + "</span>")
            cmp_ = "<br>".join(f"{e(r['title'])}: <b>{STATUS_RU[r['status']]}</b>"
                               + (f" — {e(r['detail'])}" if r.get("detail") else "") for r in d.get("compare") or [])
            extra = (f"provider: {e(d.get('provider', '') or '—')}<br>аккаунт: {e(d.get('account', '') or '—')}<br>"
                     f"смена админа: {e(d.get('last_admin_change', '') or '—')}<br>страна: {e(sd.get('country', '') or '—')}")
        else:
            w = d.get("whois") or {}
            data = e(d.get("error", ""))
            cmp_ = ""
            extra = ""
            if w:
                extra = e(f"WHOIS: {w.get('status', '')} {w.get('registrar', '')} {w.get('state', '')} "
                          f"{w.get('paid_till', '')} {w.get('error', '')}".strip())
        out.append(f'<tr class="{cls}"><td>{e(d.get("domain", ""))}</td><td>{st.get(d.get("status"), "")}</td>'
                   f"<td>{data}</td><td>{cmp_}</td><td>{extra}</td></tr>")
    return ("<table><tr><th>Домен</th><th>manager</th><th>Администратор (Sd)</th><th>Сверка</th><th>Прочее</th></tr>"
            + "".join(out) + "</table>")


def save_html(path: Path, case_id: str, app: Extraction, pas: Extraction | None, checks: list[Check],
              decision=None, case_title: str = "", domains: list[dict] | None = None):
    e = html.escape
    rows = []
    for c in checks:
        rows.append(f'<tr class="{c.status}"><td>{e(c.name)}</td><td><b>{STATUS_RU[c.status]}</b></td>'
                    f'<td>{e(c.application)}{"<br>" + _img_tag(c.crop, 70) if c.crop and c.status != "ok" else ""}</td>'
                    f'<td>{e(c.passport)}</td><td>{e(c.detail)}</td></tr>')

    def fields_table(ex: Extraction) -> str:
        out = []
        for k, f in ex.fields.items():
            v = ", ".join(f.value) if isinstance(f.value, list) else str(f.value)
            flag = ' class="review"' if f.needs_review else ""
            out.append(f"<tr{flag}><td>{e(field_title(k, ex))}</td><td>{e(v)}</td>"
                       f"<td>{e(f.source)}</td><td>{f.confidence:.2f}</td><td>{_img_tag(f.crop)}</td></tr>")
        notes = "".join(f"<li>{e(n)}</li>" for n in ex.notes)
        return (f"<p class=src>Файл: {e(ex.source_file)}</p><ul class=notes>{notes}</ul>"
                "<table><tr><th>Поле</th><th>Значение</th><th>Источник</th><th>Увер.</th><th>Фрагмент</th></tr>"
                + "".join(out) + "</table>")

    if decision is not None:
        v = decision.title + " — " + "; ".join(decision.reasons)
        vcls = {"reject": "fail", "accept": "ok", "review": "warn"}[decision.code]
    else:
        v = verdict(checks)
        vcls = "fail" if v.startswith("Есть") else "ok" if v.startswith("Расхождений") else "warn"
    doc = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Проверка {e(case_id)}</title>
<style>
body{{font-family:system-ui,Segoe UI,Arial,sans-serif;margin:24px;color:#1d1d1f;background:#fff}}
h1{{font-size:22px;margin:0 0 4px}} h2{{font-size:17px;margin:28px 0 8px}}
.verdict{{display:inline-block;padding:6px 12px;border-radius:6px;font-weight:600;margin:8px 0}}
table{{border-collapse:collapse;width:100%;font-size:14px}}
td,th{{border:1px solid #d0d0d5;padding:6px 8px;vertical-align:top;text-align:left}}
th{{background:#f3f3f6}}
.ok{{background:#e6f4ea}} .warn{{background:#fff4ce}} .fail{{background:#fde2e1}}
.review{{background:#e3eefb}} .info{{background:#f3f3f3}}
.src{{color:#666;font-size:13px}} .notes{{color:#7a4b00;font-size:13px}}
tr.grp td{{background:#f3f3f6;font-weight:600;font-size:13px}}
</style></head><body>
<h1>Проверка заявления: {e(case_id)}</h1>
<div class="src">{e(case_title)}</div>
<div class="src">Сформировано {datetime.now():%d.%m.%Y %H:%M}. Все данные обработаны локально.</div>
<div class="verdict {vcls}">{e(v)}</div>
<h2>Сверка</h2>
<table><tr><th>Проверка</th><th>Статус</th><th>Заявление</th><th>Паспорт</th><th>Комментарий</th></tr>
{''.join(rows)}</table>
{('<h2>Домены в manager</h2>' + domains_table(domains)) if domains else ''}
<h2>Заявление</h2>{fields_table(app)}
{('<h2>Паспорт</h2>' + fields_table(pas)) if pas else ''}
</body></html>"""
    path.write_text(doc, encoding="utf-8")
